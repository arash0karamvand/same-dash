"""تست آپلود اکسل حسابداری."""

from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.test import TestCase

from backend.models import Account, JournalEntry, JournalLine
from logic.accounting_excel_import import (
    DETAIL_LEDGER_SAMPLE_NOTE,
    import_excel_file,
    parse_account_code,
    parse_excel_file,
)

SAMPLE_FINANCIAL_REPORT = (
    Path(__file__).resolve().parents[1] / "agents" / "گزارشات مالی تا31 تیر.xlsx"
)


def _trial_header_rows():
    return [
        ["", None, None, None, None, None, None, ""],
        ["از تاريخ :    1405/01/01", None, None, None, None, None, None, "از سند :    1/0"],
        ["تا تاريخ :    1405/04/31", None, None, None, None, None, None, "تا سند :    100/0"],
        ["", None, None, None, None, None, None, ""],
        ["", None, None, None, None, None, None, ""],
        ["کد حساب", "عنوان حساب", "افتتاحیه", None, "گردش", None, "مانده", None],
        [None, None, "بدهکار", "بستانکار", "بدهکار", "بستانکار", "بدهکار", "بستانکار"],
    ]


def _build_sample_workbook(*, balanced=True):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    turnover_debit = 1000
    turnover_credit = 1000 if balanced else 900

    for title, rows in (
        (
            "تراز کل",
            [
                *_trial_header_rows(),
                ["7240", "فروش محصول", 0, 0, turnover_debit, turnover_credit, 0, turnover_credit],
                ["1210", "بانک", 0, 0, turnover_credit, turnover_debit, turnover_credit, 0],
                ["", "جمع", 0, 0, turnover_debit, turnover_credit, turnover_debit, turnover_credit],
            ],
        ),
        (
            "تراز معین",
            [
                *_trial_header_rows(),
                ["7240/1", "فروش", 0, 0, turnover_debit, turnover_credit, 0, turnover_credit],
                ["1210/1", "بانک", 0, 0, turnover_credit, turnover_debit, turnover_credit, 0],
                ["", "جمع", 0, 0, turnover_debit, turnover_credit, turnover_debit, turnover_credit],
            ],
        ),
        (
            "تراز تفصیلی",
            [
                *_trial_header_rows(),
                ["7240/1/1", "فروش", 0, 0, turnover_debit, turnover_credit, 0, turnover_credit],
                ["1210/1/6", "بانک ملی", 0, 0, turnover_credit, turnover_debit, turnover_credit, 0],
                ["", "جمع", 0, 0, turnover_debit, turnover_credit, turnover_debit, turnover_credit],
            ],
        ),
        (
            "ریز نمونه",
            [
                ["عنوان حساب کل :   بانک"],
                ["عنوان حساب معين :   بانک"],
                ["کد حساب تفصيلی :   1210/1/6      -      عنوان حساب تفصيلی :    بانک ملی"],
                [""],
                ["تاريخ", "شماره سند", "ع", "شرح", "بدهکار", "بستانکار", "مانده", "تش"],
                ["1405/01/01", 1, "", "افتتاحیه", 500, "", 500, "بد"],
                ["1405/01/02", 2, "1", "واریز", "", 500, 0, "بس"],
            ],
        ),
    ):
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


class AccountingExcelImportTest(TestCase):
    def test_parse_sample_workbook(self):
        parsed = parse_excel_file(_build_sample_workbook())
        self.assertEqual(len(parsed.errors), 0)
        self.assertEqual(len(parsed.general_rows), 2)
        self.assertEqual(len(parsed.subsidiary_rows), 2)
        self.assertEqual(len(parsed.detailed_rows), 2)
        self.assertEqual(len(parsed.detail_ledger_rows), 2)
        self.assertEqual(parsed.detail_ledger_account_code, "1210/1/6")
        self.assertEqual(parsed.metadata.get("date_from"), "1405/01/01")

    def test_dry_run_does_not_persist(self):
        before_accounts = Account.objects.count()
        before_entries = JournalLine.objects.count()
        report = import_excel_file(_build_sample_workbook(), dry_run=True)
        self.assertTrue(report["dry_run"])
        self.assertFalse(report["committed"])
        self.assertEqual(Account.objects.count(), before_accounts)
        self.assertEqual(JournalLine.objects.count(), before_entries)
        self.assertEqual(report["counts"]["general_rows"], 2)
        self.assertEqual(report["account_detection"]["total"], 6)
        self.assertEqual(report["account_detection"]["new"], 6)
        self.assertEqual(
            {row["level"] for row in report["recognized_accounts"]},
            {"general", "subsidiary", "detailed"},
        )

    def test_account_codes_accept_persian_digits_and_common_separators(self):
        self.assertEqual(
            parse_account_code("۱۲۱۰-۱-۶"),
            ("detailed", "1210", "1", "6"),
        )
        self.assertEqual(
            parse_account_code("۱۲۱۰.۱"),
            ("subsidiary", "1210", "1", None),
        )

    def _turnover(self, path):
        lines = JournalLine.objects.filter(
            account__path=path, journal__status_ref_id=JournalEntry.STATUS_POSTED,
        )
        return (
            sum((line.debit for line in lines), Decimal(0)),
            sum((line.credit for line in lines), Decimal(0)),
        )

    def test_import_posts_trial_balance_turnover(self):
        report = import_excel_file(_build_sample_workbook(), dry_run=False, approve=True)
        self.assertTrue(report["committed"], report["errors"])
        subsidiary = Account.objects.get(code="1", parent__code="1210")
        self.assertTrue(Account.objects.filter(code="6", parent=subsidiary).exists())
        self.assertEqual(JournalEntry.objects.count(), 1)
        journal = JournalEntry.objects.get(document_code="XL-TB-14050101-14050431")
        self.assertEqual(journal.status, JournalEntry.STATUS_POSTED)
        self.assertEqual(journal.entry_type, "opening")
        self.assertEqual(self._turnover("7240/1/1"), (Decimal(1000), Decimal(1000)))
        self.assertEqual(self._turnover("1210/1/6"), (Decimal(1000), Decimal(1000)))
        self.assertEqual(report["stats"]["journals_created"], 1)
        self.assertEqual(report["stats"]["imbalance"], 0)

    def _edited_workbook(self, edit):
        from openpyxl import load_workbook

        wb = load_workbook(_build_sample_workbook())
        edit(wb)
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def test_unbalanced_workbook_posts_difference_to_existing_retained_earnings(self):
        def edit(wb):
            wb["تراز تفصیلی"].cell(row=9, column=6, value=1100)  # 1210/1/6 credit
            ws = wb["تراز کل"]
            ws.insert_rows(10)
            for col, value in enumerate(
                ["6320", "سود و زیان انباشته", 0, 0, 0, 0, 0, 0], 1
            ):
                ws.cell(row=10, column=col, value=value)

        report = import_excel_file(self._edited_workbook(edit), dry_run=False, approve=True)
        self.assertTrue(report["committed"], report["errors"])
        self.assertEqual(report["stats"]["imbalance"], -100)
        retained = Account.objects.get(slug="retained_earnings")
        self.assertIsNone(retained.parent)
        self.assertFalse(retained.children.exists())
        self.assertEqual(self._turnover(retained.path), (Decimal(100), Decimal(0)))
        journal = JournalEntry.objects.get(document_code="XL-TB-14050101-14050431")
        self.assertEqual(journal.status, JournalEntry.STATUS_POSTED)

    def test_unbalanced_file_with_retained_earnings_children_only_warns(self):
        def edit(wb):
            general = wb["تراز کل"]
            general.insert_rows(10)
            for col, value in enumerate(
                ["6320", "سود و زیان انباشته", 0, 0, 0, 0, 0, 0], 1
            ):
                general.cell(row=10, column=col, value=value)
            subsidiary = wb["تراز معین"]
            subsidiary.insert_rows(10)
            for col, value in enumerate(
                ["6320/1", "سنواتی", 0, 0, 0, 0, 0, 0], 1
            ):
                subsidiary.cell(row=10, column=col, value=value)
            detailed = wb["تراز تفصیلی"]
            detailed.cell(row=9, column=6, value=1100)

        report = import_excel_file(
            self._edited_workbook(edit),
            dry_run=False,
            approve=False,
        )
        self.assertTrue(report["committed"], report["errors"])
        self.assertEqual(report["errors"], [])
        self.assertNotEqual(report["stats"]["imbalance"], 0)
        balancing = Account.objects.get(slug="retained_earnings-import-diff")
        self.assertEqual(balancing.parent.code, "6320")
        self.assertTrue(balancing.is_postable)
        self.assertTrue(any("فایل ناتراز است" in warning for warning in report["warnings"]))

    def test_detailed_rows_win_when_they_cover_the_general_total(self):
        def edit(wb):
            # معین بدون تفصیلی که مبلغش در تفصیلیِ معین دیگر آمده است
            ws = wb["تراز معین"]
            ws.insert_rows(8)
            for col, value in enumerate(["7240/2", "فروش ۲", 0, 0, 0, 300, 0, 300], 1):
                ws.cell(row=8, column=col, value=value)

        report = import_excel_file(self._edited_workbook(edit), dry_run=False, approve=True)
        self.assertTrue(report["committed"], report["errors"])
        self.assertEqual(self._turnover("7240/2"), (Decimal(0), Decimal(0)))
        self.assertEqual(report["stats"]["imbalance"], 0)
        self.assertTrue(any("7240/2" in warning for warning in report["warnings"]))

    def test_duplicate_import_requires_replace(self):
        def net(path):
            debit, credit = self._turnover(path)
            return debit - credit

        workbook = lambda: _build_sample_workbook(balanced=False)  # noqa: E731
        import_excel_file(workbook(), dry_run=False, approve=True)
        second = import_excel_file(workbook(), dry_run=False, approve=True)
        self.assertEqual(second["stats"]["journals_created"], 0)
        self.assertEqual(net("7240/1/1"), Decimal(100))

        third = import_excel_file(workbook(), dry_run=False, approve=True, force=True)
        self.assertEqual(third["stats"]["journals_created"], 1)
        self.assertTrue(JournalEntry.objects.filter(document_code="XL-TB-14050101-14050431-R2").exists())
        self.assertTrue(JournalEntry.objects.filter(corrects__document_code="XL-TB-14050101-14050431").exists())
        self.assertEqual(net("7240/1/1"), Decimal(100))

    def test_new_period_reuses_accounts_and_adds_turnover(self):
        first = import_excel_file(_build_sample_workbook(), dry_run=False, approve=True)
        self.assertTrue(first["committed"], first["errors"])

        def edit(wb):
            for sheet in ("تراز کل", "تراز معین", "تراز تفصیلی"):
                ws = wb[sheet]
                ws.cell(row=2, column=1, value="از تاريخ :    1405/05/01")
                ws.cell(row=3, column=1, value="تا تاريخ :    1405/05/31")

        second = import_excel_file(self._edited_workbook(edit), dry_run=False, approve=True)
        self.assertTrue(second["committed"], second["errors"])
        self.assertEqual(second["stats"]["accounts_created"], 0)
        self.assertEqual(second["stats"]["subsidiaries_created"], 0)
        self.assertEqual(second["stats"]["details_created"], 0)
        self.assertEqual(second["stats"]["journals_created"], 1)
        self.assertTrue(
            JournalEntry.objects.filter(document_code="XL-TB-14050501-14050531").exists()
        )
        self.assertEqual(self._turnover("1210/1/6"), (Decimal(2000), Decimal(2000)))

    def test_financial_report_sample_workbook(self):
        if not SAMPLE_FINANCIAL_REPORT.is_file():
            self.skipTest("sample workbook missing")
        with SAMPLE_FINANCIAL_REPORT.open("rb") as handle:
            parsed = parse_excel_file(handle)
        self.assertEqual(parsed.errors, [])
        self.assertEqual(len(parsed.general_rows), 32)
        self.assertEqual(len(parsed.subsidiary_rows), 142)
        self.assertGreaterEqual(len(parsed.detailed_rows), 1200)
        self.assertEqual(parsed.detail_ledger_account_code, "1210/1/6")
        self.assertEqual(parsed.metadata.get("date_from"), "1405/01/01")
        self.assertEqual(parsed.metadata.get("date_to"), "1405/04/31")

        with SAMPLE_FINANCIAL_REPORT.open("rb") as handle:
            report = import_excel_file(handle, dry_run=True)
        self.assertFalse(report["errors"])
        self.assertTrue(report["import_mode"]["chart_from_trial_balance"])
        self.assertIn(DETAIL_LEDGER_SAMPLE_NOTE, report["warnings"])
        self.assertGreater(report["stats"]["subsidiaries_created"], 0)
        self.assertGreater(report["stats"]["details_created"], 0)
        self.assertGreater(report["stats"]["turnover_lines"], 1500)
        self.assertEqual(report["stats"]["imbalance"], 11345226)
