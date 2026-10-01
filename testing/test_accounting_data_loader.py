"""تست خواندن گزارش چاپی اکسل حسابداری."""

import os
import tempfile
import unittest
from pathlib import Path

import pandas as pd

from openpyxl import Workbook

from logic.accounting_data_loader import (
    RIZ_COLUMNS,
    RIZ_MONEY_COLUMNS,
    TARAZ_COLUMNS,
    AccountingDataLoader,
    balances_from_records,
    movement_records,
)

SAMPLE_FINANCIAL_REPORT = (
    Path(__file__).resolve().parents[1] / "agents" / "گزارشات مالی تا31 تیر.xlsx"
)


def _save_workbook(workbook: Workbook) -> str:
    handle, name = tempfile.mkstemp(suffix=".xlsx")
    os.close(handle)
    workbook.save(name)
    return name


class AccountingDataLoaderTests(unittest.TestCase):
    def test_balance_comes_from_records_not_the_reported_column(self):
        trial = pd.DataFrame([
            {
                "کد حساب": "1120",
                "عنوان حساب": "صندوق",
                "افتتاحیه بدهکار": 0,
                "افتتاحیه بستانکار": 0,
                "گردش بدهکار": 10,
                "گردش بستانکار": 4,
                "مانده بدهکار": 999,
                "مانده بستانکار": 0,
            }
        ])
        records = movement_records(trial)
        balances = balances_from_records(records)
        self.assertEqual(len(records), 2)
        self.assertEqual(set(records["نوع رکورد"]), {"گردش"})
        self.assertEqual(int(balances.loc[0, "مانده بدهکار"]), 6)
        self.assertEqual(int(balances.loc[0, "مانده بستانکار"]), 0)

    def test_printed_trial_balance_layout_is_cleaned(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "تراز كل"
        for _ in range(4):
            sheet.append(["سربرگ", None, None, None, None, None, None, None])
        sheet.append(["کد حساب", "عنوان حساب", "افتتاحیه", None, "گردش", None, "مانده", None])
        sheet.append([None, None, "بدهکار", "بستانکار", "بدهکار", "بستانکار", "بدهکار", "بستانکار"])
        sheet.append([1120, "صندوق", None, None, "1,500", 400, 1100, None])
        sheet.append([1300.0, None, 5, 0, 0, 0, 5, 0])
        sheet.append([None, None, None, None, None, None, None, None])
        sheet.append([None, "جمع", None, None, 1500, 400, 1100, None])
        path = _save_workbook(workbook)
        loader = AccountingDataLoader(path)
        try:
            frame = loader.load_taraz("تراز کل")
            self.assertEqual(loader.errors, [])
            self.assertEqual(list(frame.columns), TARAZ_COLUMNS)
            self.assertEqual(len(frame), 1)
            self.assertEqual(frame.loc[0, "کد حساب"], "1120")
            self.assertEqual(frame.loc[0, "افتتاحیه بدهکار"], 0)
            self.assertEqual(frame.loc[0, "گردش بدهکار"], 1500)
            self.assertEqual(frame.loc[0, "مانده بدهکار"], 1100)
            self.assertEqual(frame.loc[0, "مانده بستانکار"], 0)
            self.assertEqual(int(frame["مانده بدهکار"].sum()), 1100)
        finally:
            loader.close()
            os.remove(path)

    def test_detail_ledger_renames_debit_and_credit(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ریز نمونه"
        sheet.append(["عنوان حساب کل : بانک"])
        sheet.append(["عنوان حساب معین : بانک"])
        sheet.append(["کد حساب تفصیلی : 1210/1/6 - عنوان حساب تفصیلی : بانک ملی"])
        sheet.append(["", None, None, None, None, None, None, None])
        sheet.append(
            [
                "تاريخ",
                "شماره سند",
                "ع",
                "شرح",
                "بدهکار - افزایش حساب",
                "بستانکار- کاهش حساب",
                "مانده",
                "تش.",
            ]
        )
        sheet.append(["1405/01/01", 3, None, "افتتاحیه", "1,500", None, 1500, "بد"])
        sheet.append(["1405/01/02", 4, None, "کارمزد", None, "۲۰۰", 1300, "بد"])
        sheet.append([None, None, None, "جمع", 1500, 200, 1300, "بد"])
        path = _save_workbook(workbook)
        loader = AccountingDataLoader(path)
        try:
            frame = loader.load_riz_nemooneh()
            self.assertEqual(loader.errors, [])
            self.assertEqual(list(frame.columns), RIZ_COLUMNS)
            self.assertEqual(len(frame), 2)
            self.assertEqual(frame.loc[0, "بدهکار"], 1500)
            self.assertEqual(frame.loc[0, "بستانکار"], 0)
            self.assertEqual(frame.loc[0, "ع"], 0)
            self.assertEqual(frame.loc[1, "بستانکار"], 200)
            self.assertEqual(frame.attrs["report"]["detailed_code"], "1210/1/6")
            self.assertNotIn("جمع", set(frame["شرح"]))
        finally:
            loader.close()
            os.remove(path)

    def test_missing_sheet_does_not_raise(self):
        workbook = Workbook()
        workbook.active.title = "تراز کل"
        path = _save_workbook(workbook)
        loader = AccountingDataLoader(path)
        missing_file = AccountingDataLoader(path + ".missing")
        try:
            frame = loader.load_taraz("تراز معین")
            empty = missing_file.load_riz_nemooneh()
            self.assertTrue(frame.empty)
            self.assertEqual(list(frame.columns), TARAZ_COLUMNS)
            self.assertTrue(loader.errors)
            self.assertTrue(empty.empty)
            self.assertEqual(list(empty.columns), RIZ_COLUMNS)
            self.assertTrue(missing_file.errors)
        finally:
            loader.close()
            missing_file.close()
            os.remove(path)

    @unittest.skipUnless(SAMPLE_FINANCIAL_REPORT.is_file(), "sample workbook missing")
    def test_real_financial_report(self):
        loader = AccountingDataLoader(SAMPLE_FINANCIAL_REPORT)
        try:
            general = loader.load_taraz("تراز کل")
            subsidiary = loader.load_taraz("تراز معین")
            detailed = loader.load_taraz("تراز تفصیلی")
            ledger = loader.load_riz_nemooneh("ریز نمونه")
        finally:
            loader.close()

        self.assertEqual(loader.errors, [])
        self.assertEqual(loader.warnings, [])
        self.assertEqual(len(general), 32)
        self.assertEqual(len(subsidiary), 142)
        self.assertEqual(len(detailed), 1226)
        self.assertEqual(len(ledger), 115)
        self.assertEqual(int(general["مانده بدهکار"].sum()), 1_734_262_965_929)
        calculated = balances_from_records(movement_records(general))
        self.assertEqual(int(calculated["مانده بدهکار"].sum()), 1_734_262_965_929)
        reported_net = general["مانده بدهکار"] - general["مانده بستانکار"]
        calculated_net = (
            general["افتتاحیه بدهکار"] - general["افتتاحیه بستانکار"]
            + general["گردش بدهکار"] - general["گردش بستانکار"]
        )
        self.assertTrue((reported_net == calculated_net).all())
        self.assertEqual(int(general["گردش بدهکار"].sum()), 4_868_289_562_186)
        self.assertEqual(int(general["گردش بستانکار"].sum()), 4_868_278_216_960)
        self.assertEqual(int(general["مانده بستانکار"].sum()), 1_734_251_620_703)
        self.assertEqual(general.attrs["report"]["date_from"], "1405/01/01")
        self.assertEqual(general.attrs["report"]["date_to"], "1405/04/31")
        self.assertEqual(ledger.attrs["report"]["detailed_code"], "1210/1/6")
        self.assertEqual(detailed.attrs["report"].get("sample_accounts"), ["1210/1/6"])
        self.assertEqual(int(ledger["بدهکار"].sum()), 29_734_110_800)
        self.assertEqual(int(ledger["بستانکار"].sum()), 28_033_044_177)
        self.assertEqual(int(ledger.iloc[-1]["مانده"]), 1_701_066_623)
        closed = general.loc[general["کد حساب"] == "1220"].iloc[0]
        self.assertEqual(closed["مانده بدهکار"], 0)
        self.assertEqual(closed["مانده بستانکار"], 0)
        self.assertFalse(general[TARAZ_COLUMNS[2:]].isna().any().any())
        self.assertFalse(ledger[RIZ_MONEY_COLUMNS].isna().any().any())


if __name__ == "__main__":
    unittest.main()
