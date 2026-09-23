"""آپلود و پردازش فایل‌های اکسل حسابداری (تراز کل/معین/تفصیلی + ریز)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction

from backend.models import Account, JournalEntry
from backend.models.accounting import defer_account_closure_rebuild
from logic.chart_of_accounts import infer_account_class, infer_normal_balance, slug_for_code
from logic.jalali import parse_jalali_date
from logic.ledger import OFFICE_LEDGER

SHEET_ALIASES = {
    "general": ("تراز کل", "تراز كل"),
    "subsidiary": ("تراز معین", "تراز معين", "تراز معي"),
    "detailed": ("تراز تفصیلی", "تراز تفصيلي", "تراز تفصيل"),
    "detail_ledger": ("ریز نمونه", "ريز نمونه", "ریز", "ريز"),
}

DETAIL_LEDGER_SAMPLE_NOTE = (
    "شیت «ریز نمونه» فقط گردش یک حساب را نشان می‌دهد و سند جداگانه‌ای از آن ساخته نمی‌شود؛ "
    "مانده و گردش همه حساب‌ها از تراز تفصیلی ثبت می‌شود."
)

TOTAL_LABELS = {"جمع", "جمع کل", "جمع كل"}
DETAILED_CODE_RE = re.compile(r"(\d{3,4}/\d+/[\d]+)")
DATE_RANGE_RE = re.compile(r"(\d{4}/\d{2}/\d{2})")
DOC_RANGE_RE = re.compile(r"(\d+)\s*/\s*(\d+)")


@dataclass
class TrialBalanceRow:
    code: str
    name: str
    opening_debit: Decimal = Decimal(0)
    opening_credit: Decimal = Decimal(0)
    turnover_debit: Decimal = Decimal(0)
    turnover_credit: Decimal = Decimal(0)
    balance_debit: Decimal = Decimal(0)
    balance_credit: Decimal = Decimal(0)


@dataclass
class DetailLedgerRow:
    entry_date: str
    document_number: int
    attach_code: str
    description: str
    debit: Decimal = Decimal(0)
    credit: Decimal = Decimal(0)


@dataclass
class ParsedExcel:
    metadata: dict = field(default_factory=dict)
    general_rows: list[TrialBalanceRow] = field(default_factory=list)
    subsidiary_rows: list[TrialBalanceRow] = field(default_factory=list)
    detailed_rows: list[TrialBalanceRow] = field(default_factory=list)
    detail_ledger_account_code: str = ""
    detail_ledger_account_name: str = ""
    detail_ledger_rows: list[DetailLedgerRow] = field(default_factory=list)
    recognized_accounts: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _normalize_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    return (
        text.replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", "")
        .translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))
        .strip()
    )


def _normalize_account_code(value):
    code = _normalize_text(value)
    if not code:
        return ""
    if re.fullmatch(r"\d+\.0+", code):
        code = code.split(".", 1)[0]
    code = code.replace("\\", "/").replace("ـ", "")
    code = re.sub(r"(?<=\d)\s*[-._]\s*(?=\d)", "/", code)
    code = re.sub(r"\s*/\s*", "/", code)
    return code.strip(" /")


def _parse_money(value):
    if value is None or value == "":
        return Decimal(0)
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = _normalize_text(value).replace(",", "")
    if not text:
        return Decimal(0)
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"مبلغ نامعتبر: {value}") from exc


def _find_sheet(workbook, aliases):
    normalized = {_normalize_text(name): name for name in workbook.sheetnames}
    for alias in aliases:
        key = _normalize_text(alias)
        if key in normalized:
            return workbook[normalized[key]]
    return None


def _parse_metadata_rows(rows):
    meta = {}
    for row in rows[:5]:
        joined = " ".join(_normalize_text(cell) for cell in row if cell is not None)
        if not joined:
            continue
        dates = DATE_RANGE_RE.findall(joined)
        if "از تاريخ" in joined.replace(" ", "") or "از تاریخ" in joined:
            if dates:
                meta.setdefault("date_from", dates[0])
        if "تا تاريخ" in joined.replace(" ", "") or "تا تاریخ" in joined:
            if dates:
                meta.setdefault("date_to", dates[-1])
        doc_match = DOC_RANGE_RE.search(joined)
        if doc_match and ("سند" in joined or "سند" in joined.replace("ي", "ی")):
            if "از سند" in joined.replace(" ", ""):
                meta.setdefault("doc_from", int(doc_match.group(1)))
            elif "تا سند" in joined.replace(" ", ""):
                meta.setdefault("doc_to", int(doc_match.group(1)))
    return meta


def _find_header_row(rows):
    for index, row in enumerate(rows):
        first = _normalize_text(row[0] if row else "")
        second = _normalize_text(row[1] if row and len(row) > 1 else "")
        if first == "کد حساب" and "عنوان" in second:
            return index
    return None


def _parse_trial_balance_sheet(worksheet):
    rows = list(worksheet.iter_rows(values_only=True))
    metadata = _parse_metadata_rows(rows)
    header_index = _find_header_row(rows)
    if header_index is None:
        raise ValueError("سطر عنوان ستون‌های تراز (کد حساب / عنوان حساب) یافت نشد.")

    parsed_rows = []
    totals = None
    data_start = header_index + 2
    for row in rows[data_start:]:
        code = _normalize_account_code(row[0] if row else "")
        name = _normalize_text(row[1] if row and len(row) > 1 else "")
        if not code and not name:
            continue
        if name in TOTAL_LABELS or code in TOTAL_LABELS:
            totals = TrialBalanceRow(
                code="",
                name="جمع",
                opening_debit=_parse_money(row[2] if row and len(row) > 2 else 0),
                opening_credit=_parse_money(row[3] if row and len(row) > 3 else 0),
                turnover_debit=_parse_money(row[4] if row and len(row) > 4 else 0),
                turnover_credit=_parse_money(row[5] if row and len(row) > 5 else 0),
                balance_debit=_parse_money(row[6] if row and len(row) > 6 else 0),
                balance_credit=_parse_money(row[7] if row and len(row) > 7 else 0),
            )
            break
        if not code:
            continue
        parsed_rows.append(
            TrialBalanceRow(
                code=code,
                name=name,
                opening_debit=_parse_money(row[2] if row and len(row) > 2 else 0),
                opening_credit=_parse_money(row[3] if row and len(row) > 3 else 0),
                turnover_debit=_parse_money(row[4] if row and len(row) > 4 else 0),
                turnover_credit=_parse_money(row[5] if row and len(row) > 5 else 0),
                balance_debit=_parse_money(row[6] if row and len(row) > 6 else 0),
                balance_credit=_parse_money(row[7] if row and len(row) > 7 else 0),
            )
        )
    return metadata, parsed_rows, totals


def _parse_detail_ledger_sheet(worksheet):
    rows = list(worksheet.iter_rows(values_only=True))
    account_code = ""
    account_name = ""
    for row in rows[:4]:
        joined = " ".join(_normalize_text(cell) for cell in row if cell is not None)
        match = DETAILED_CODE_RE.search(joined)
        if match:
            account_code = match.group(1)
        if "عنوان حساب تفصيلی" in joined or "عنوان حساب تفصیلی" in joined:
            parts = joined.split(":")
            if len(parts) > 1:
                account_name = parts[-1].strip(" -")

    header_index = None
    for index, row in enumerate(rows):
        first = _normalize_text(row[0] if row else "")
        if first in ("تاريخ", "تاریخ"):
            header_index = index
            break
    if header_index is None:
        raise ValueError("سطر عنوان ستون‌های ریز حساب یافت نشد.")

    parsed_rows = []
    for row in rows[header_index + 1 :]:
        date_text = _normalize_text(row[0] if row else "")
        if not date_text:
            continue
        doc_raw = row[1] if row and len(row) > 1 else None
        if doc_raw is None or str(doc_raw).strip() == "":
            continue
        try:
            document_number = int(doc_raw)
        except (TypeError, ValueError):
            continue
        parsed_rows.append(
            DetailLedgerRow(
                entry_date=date_text,
                document_number=document_number,
                attach_code=_normalize_text(row[2] if row and len(row) > 2 else ""),
                description=_normalize_text(row[3] if row and len(row) > 3 else ""),
                debit=_parse_money(row[4] if row and len(row) > 4 else 0),
                credit=_parse_money(row[5] if row and len(row) > 5 else 0),
            )
        )
    return account_code, account_name, parsed_rows


def parse_account_code(code):
    parts = [part.strip() for part in _normalize_account_code(code).split("/") if part.strip()]
    if not parts or any(not part.isdigit() for part in parts):
        raise ValueError(f"کد حساب نامعتبر: {code}")
    if len(parts) == 1:
        return "general", parts[0], None, None
    if len(parts) == 2:
        return "subsidiary", parts[0], parts[1], None
    if len(parts) >= 3:
        return "detailed", parts[0], parts[1], parts[2]
    raise ValueError(f"کد حساب نامعتبر: {code}")


def _recognized_accounts(parsed, *, ledger=OFFICE_LEDGER):
    """فهرست حساب‌هایی که واقعاً از فایل تشخیص داده شده‌اند."""
    existing = set(ledger.accounts().values_list("path", flat=True))
    recognized = []
    seen = {}
    conflicts = []
    for level, rows in (
        ("general", parsed.general_rows),
        ("subsidiary", parsed.subsidiary_rows),
        ("detailed", parsed.detailed_rows),
    ):
        for row in rows:
            try:
                parsed_level, general, subsidiary, detail = parse_account_code(row.code)
            except ValueError as exc:
                parsed.errors.append(str(exc))
                continue
            if parsed_level != level:
                parsed.warnings.append(
                    f"سطح کد {row.code} با شیت {level} هم‌خوان نیست."
                )
                continue
            path = "/".join(part for part in (general, subsidiary, detail) if part)
            previous = seen.get(path)
            if previous and previous != row.name:
                conflicts.append(
                    f"کد {path} با دو عنوان «{previous}» و «{row.name}» دیده شد."
                )
                continue
            if previous:
                continue
            seen[path] = row.name
            recognized.append({
                "level": level,
                "code": row.code,
                "path": path,
                "name": row.name or f"حساب {path}",
                "parent_path": path.rpartition("/")[0],
                "status": "existing" if path in existing else "new",
            })
    parsed.errors.extend(conflicts)
    return recognized


def validate_excel_workbook(workbook):
    parsed = ParsedExcel()
    general_ws = _find_sheet(workbook, SHEET_ALIASES["general"])
    subsidiary_ws = _find_sheet(workbook, SHEET_ALIASES["subsidiary"])
    detailed_ws = _find_sheet(workbook, SHEET_ALIASES["detailed"])
    detail_ws = _find_sheet(workbook, SHEET_ALIASES["detail_ledger"])

    if not general_ws:
        parsed.errors.append("شیت «تراز کل» یافت نشد.")
        return parsed
    if not subsidiary_ws:
        parsed.errors.append("شیت «تراز معین» یافت نشد.")
        return parsed
    if not detailed_ws:
        parsed.errors.append("شیت «تراز تفصیلی» یافت نشد.")
        return parsed

    try:
        general_meta, parsed.general_rows, general_totals = _parse_trial_balance_sheet(general_ws)
        parsed.metadata.update(general_meta)
        parsed.metadata["general_totals"] = general_totals
    except ValueError as exc:
        parsed.errors.append(str(exc))

    try:
        subsidiary_meta, parsed.subsidiary_rows, subsidiary_totals = _parse_trial_balance_sheet(subsidiary_ws)
        parsed.metadata.setdefault("date_from", subsidiary_meta.get("date_from"))
        parsed.metadata.setdefault("date_to", subsidiary_meta.get("date_to"))
        parsed.metadata["subsidiary_totals"] = subsidiary_totals
    except ValueError as exc:
        parsed.errors.append(str(exc))

    try:
        _, parsed.detailed_rows, detailed_totals = _parse_trial_balance_sheet(detailed_ws)
        parsed.metadata["detailed_totals"] = detailed_totals
    except ValueError as exc:
        parsed.errors.append(str(exc))

    if detail_ws:
        try:
            (
                parsed.detail_ledger_account_code,
                parsed.detail_ledger_account_name,
                parsed.detail_ledger_rows,
            ) = _parse_detail_ledger_sheet(detail_ws)
        except ValueError as exc:
            parsed.errors.append(str(exc))

    for label, totals in (
        ("تراز کل", parsed.metadata.get("general_totals")),
        ("تراز معین", parsed.metadata.get("subsidiary_totals")),
        ("تراز تفصیلی", parsed.metadata.get("detailed_totals")),
    ):
        if not totals:
            continue
        if totals.turnover_debit != totals.turnover_credit:
            parsed.warnings.append(
                f"{label}: گردش بدهکار ({int(totals.turnover_debit)}) "
                f"با بستانکار ({int(totals.turnover_credit)}) برابر نیست."
            )
        if totals.opening_debit != totals.opening_credit:
            parsed.warnings.append(
                f"{label}: افتتاحیه بدهکار ({int(totals.opening_debit)}) "
                f"با بستانکار ({int(totals.opening_credit)}) برابر نیست."
            )

    if parsed.detail_ledger_rows:
        detail_debit = sum(row.debit for row in parsed.detail_ledger_rows)
        detail_credit = sum(row.credit for row in parsed.detail_ledger_rows)
        if detail_debit != detail_credit:
            parsed.warnings.append(
                f"ریز حساب (تک‌حسابی): جمع بدهکار ({int(detail_debit)}) "
                f"و بستانکار ({int(detail_credit)}) — این طبیعی است و مانع import نیست."
            )

    return parsed


def parse_excel_file(file_obj):
    from openpyxl import load_workbook

    content = file_obj.read()
    if not content:
        raise ValueError("فایل خالی است.")
    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    try:
        parsed = validate_excel_workbook(workbook)
    finally:
        workbook.close()
    return parsed


def _get_or_create_general_account(code, name, *, ledger=OFFICE_LEDGER):
    account = ledger.accounts().filter(code=code, parent__isnull=True).first()
    if account:
        clean_name = name.strip()
        if clean_name and account.name != clean_name:
            account.name = clean_name
            account.save(update_fields=["name"])
        return account, False

    slug = slug_for_code(code, ledger=ledger)
    account_class = infer_account_class(code)
    account, created = Account.objects.get_or_create(
        ledger=ledger.model, slug=slug,
        defaults={
            "code": code,
            "name": name.strip() or f"حساب {code}",
            "account_class": account_class,
            "normal_balance": infer_normal_balance(account_class),
            "sort_order": int(code) if code.isdigit() else 0,
            "is_active": True,
        },
    )
    if not created and not account.code:
        account.code = code
        account.save(update_fields=["code"])
    return account, created


def _get_or_create_subsidiary(general_code, sub_code, name, *, ledger=OFFICE_LEDGER):
    account = ledger.accounts().filter(code=general_code, parent__isnull=True).first()
    if not account:
        account, _ = _get_or_create_general_account(general_code, name, ledger=ledger)
    sub = ledger.accounts().filter(parent=account, code=sub_code).first()
    if sub:
        clean_name = name.strip()
        if clean_name and sub.name != clean_name:
            sub.name = clean_name
            sub.save(update_fields=["name"])
        return sub, False
    sub = Account.objects.create(
        ledger=ledger.model, parent=account, slug=f"{account.slug}-{sub_code}",
        code=sub_code,
        name=name.strip() or f"معین {general_code}/{sub_code}",
        account_class=account.account_class, normal_balance=account.normal_balance,
    )
    return sub, True


def _get_or_create_detailed(general_code, sub_code, detail_code, name, *, ledger=OFFICE_LEDGER):
    subsidiary = ledger.accounts().filter(parent__code=general_code, code=sub_code).first()
    if not subsidiary:
        subsidiary, _ = _get_or_create_subsidiary(general_code, sub_code, name, ledger=ledger)
    detail = ledger.accounts().filter(parent=subsidiary, code=detail_code).first()
    if detail:
        clean_name = name.strip()
        if clean_name and detail.name != clean_name:
            detail.name = clean_name
            detail.save(update_fields=["name"])
        return detail, False
    detail = Account.objects.create(
        ledger=ledger.model, parent=subsidiary, slug=f"{subsidiary.slug}-{detail_code}",
        code=detail_code,
        name=name.strip() or f"تفصیلی {general_code}/{sub_code}/{detail_code}",
        account_class=subsidiary.account_class, normal_balance=subsidiary.normal_balance,
    )
    return detail, True


def _estimate_import_stats(parsed, *, ledger=OFFICE_LEDGER):
    """Read-only preview — no writes."""
    ledger_row = ledger.model
    general_existing = set(
        Account.objects.filter(ledger=ledger_row, parent__isnull=True).values_list("code", flat=True)
    )
    sub_existing = set(
        Account.objects.filter(
            ledger=ledger_row,
            parent__isnull=False,
            parent__parent__isnull=True,
        ).values_list("parent__code", "code")
    )
    detail_existing = set(
        Account.objects.filter(
            ledger=ledger_row,
            parent__parent__isnull=False,
            parent__parent__parent__isnull=True,
        ).values_list("parent__parent__code", "parent__code", "code")
    )

    accounts_created = accounts_updated = 0
    for row in parsed.general_rows:
        if row.code in general_existing:
            accounts_updated += 1
        else:
            accounts_created += 1

    subsidiaries_created = subsidiaries_updated = 0
    for row in parsed.subsidiary_rows:
        level, general_code, sub_code, _ = parse_account_code(row.code)
        if level != "subsidiary":
            continue
        if (general_code, sub_code) in sub_existing:
            subsidiaries_updated += 1
        else:
            subsidiaries_created += 1

    details_created = details_updated = 0
    for row in parsed.detailed_rows:
        level, general_code, sub_code, detail_code = parse_account_code(row.code)
        if level != "detailed":
            continue
        if (general_code, sub_code, detail_code) in detail_existing:
            details_updated += 1
        else:
            details_created += 1

    return {
        "accounts_created": accounts_created,
        "accounts_updated": accounts_updated,
        "subsidiaries_created": subsidiaries_created,
        "subsidiaries_updated": subsidiaries_updated,
        "details_created": details_created,
        "details_updated": details_updated,
        **_estimate_balance_stats(parsed, ledger=ledger),
    }


def _estimate_balance_stats(parsed, *, ledger):
    opening, turnover = _balance_lines(_leaf_balance_rows(parsed))
    codes = [code for code, lines in zip(_balance_document_codes(parsed), (opening, turnover)) if lines]
    return {
        "opening_lines": len(opening),
        "turnover_lines": len(turnover),
        "imbalance": int(_imbalance(opening) + _imbalance(turnover)),
        "journals_created": len(codes),
        "journals_existing": sum(
            _live_import_journals(code, ledger=ledger).count() for code in codes
        ),
    }


def _import_chart_rows(parsed, *, ledger=OFFICE_LEDGER):
    stats = {
        "accounts_created": 0,
        "accounts_updated": 0,
        "subsidiaries_created": 0,
        "subsidiaries_updated": 0,
        "details_created": 0,
        "details_updated": 0,
    }

    for row in parsed.general_rows:
        _, created = _get_or_create_general_account(row.code, row.name, ledger=ledger)
        if created:
            stats["accounts_created"] += 1
        else:
            stats["accounts_updated"] += 1

    for row in parsed.subsidiary_rows:
        level, general_code, sub_code, _ = parse_account_code(row.code)
        if level != "subsidiary":
            parsed.warnings.append(f"ردیف معین نادیده گرفته شد: {row.code}")
            continue
        _, created = _get_or_create_subsidiary(general_code, sub_code, row.name, ledger=ledger)
        if created:
            stats["subsidiaries_created"] += 1
        else:
            stats["subsidiaries_updated"] += 1

    for row in parsed.detailed_rows:
        level, general_code, sub_code, detail_code = parse_account_code(row.code)
        if level != "detailed":
            parsed.warnings.append(f"ردیف تفصیلی نادیده گرفته شد: {row.code}")
            continue
        _, created = _get_or_create_detailed(general_code, sub_code, detail_code, row.name, ledger=ledger)
        if created:
            stats["details_created"] += 1
        else:
            stats["details_updated"] += 1

    return stats


def _leaf_balance_rows(parsed):
    """ردیف‌های قابل ثبت هر حساب کل.

    اگر جمع تفصیلی‌های یک کل با تراز کل برابر باشد، فقط تفصیلی‌ها ثبت می‌شوند
    (خروجی برخی نرم‌افزارها معینِ بدون تفصیلی را در تفصیلیِ معین دیگری می‌آورد).
    وگرنه هر معین با تفصیلی‌هایش، و معینِ بدون تفصیلی با مبلغ خودش ثبت می‌شود.
    """
    details, subs = {}, {}
    for row in parsed.detailed_rows:
        level, general_code, sub_code, _ = parse_account_code(row.code)
        if level == "detailed":
            details.setdefault(general_code, {}).setdefault(sub_code, []).append(row)
    for row in parsed.subsidiary_rows:
        level, general_code, sub_code, _ = parse_account_code(row.code)
        if level == "subsidiary":
            subs.setdefault(general_code, {})[sub_code] = row

    leaves = []
    for general in parsed.general_rows:
        general_details = details.get(general.code, {})
        general_subs = subs.get(general.code, {})
        detail_rows = [row for rows in general_details.values() for row in rows]
        if not general_subs and not detail_rows:
            leaves.append(general)
            continue
        mismatched = [
            sub.code for sub_code, sub in general_subs.items()
            if _row_amounts(sub) != _sum_amounts(general_details.get(sub_code, []))
        ]
        if detail_rows and _sum_amounts(detail_rows) == _row_amounts(general):
            leaves.extend(detail_rows)
        else:
            leaves.extend(detail_rows)
            leaves.extend(
                sub for sub_code, sub in general_subs.items() if sub_code not in general_details
            )
            mismatched = [code for code in mismatched if parse_account_code(code)[2] in general_details]
        if mismatched:
            parsed.warnings.append(
                f"حساب {general.code}: مانده معین‌های {'، '.join(mismatched[:6])}"
                f"{' …' if len(mismatched) > 6 else ''} با جمع تفصیلی‌هایشان برابر نیست؛ "
                "مبالغ تفصیلی ثبت شد."
            )
    return leaves


def _row_amounts(row):
    return (
        row.opening_debit - row.opening_credit,
        row.turnover_debit,
        row.turnover_credit,
    )


def _sum_amounts(rows):
    return (
        sum((r.opening_debit - r.opening_credit for r in rows), Decimal(0)),
        sum((r.turnover_debit for r in rows), Decimal(0)),
        sum((r.turnover_credit for r in rows), Decimal(0)),
    )


def _balance_lines(leaves):
    """ردیف‌های یک‌طرفه برای سند افتتاحیه و سند گردش دوره."""
    opening, turnover = [], []
    for row in leaves:
        net_opening = row.opening_debit - row.opening_credit
        if net_opening > 0:
            opening.append((row, net_opening, Decimal(0)))
        elif net_opening < 0:
            opening.append((row, Decimal(0), -net_opening))
        if row.turnover_debit > 0:
            turnover.append((row, row.turnover_debit, Decimal(0)))
        if row.turnover_credit > 0:
            turnover.append((row, Decimal(0), row.turnover_credit))
    return opening, turnover


def _imbalance(lines):
    return sum((d for _, d, _ in lines), Decimal(0)) - sum((c for _, _, c in lines), Decimal(0))


def _balance_document_codes(parsed):
    date_from = (parsed.metadata.get("date_from") or "").replace("/", "")
    date_to = (parsed.metadata.get("date_to") or "").replace("/", "")
    return f"XL-OB-{date_from}", f"XL-TB-{date_from}-{date_to}"


def _balancing_account(ledger):
    """حساب فنی اختلاف؛ ناترازی فایل هرگز مانع ورود اطلاعات نمی‌شود."""
    from logic.chart_of_accounts import ACCOUNT_SLUGS

    general = ledger.accounts().filter(slug=ACCOUNT_SLUGS.RETAINED_EARNINGS, parent__isnull=True).first()
    if general is None:
        general = Account.objects.create(
            ledger=ledger.model,
            parent=None,
            slug=ACCOUNT_SLUGS.RETAINED_EARNINGS,
            code="6320",
            name="سود و زیان انباشته",
            account_class="equity",
            normal_balance="credit",
            sort_order=6320,
            is_active=True,
        )
    if not general.children.exists():
        return general
    slug = f"{general.slug}-import-diff"
    existing = general.children.filter(slug=slug).first()
    if existing is not None:
        if not existing.is_active:
            existing.is_active = True
            existing.save(update_fields=["is_active"])
        if existing.is_postable:
            return existing

    code = "9999"
    suffix = 1
    while general.children.filter(code=code).exists():
        suffix += 1
        code = f"9999{suffix}"
    return Account.objects.create(
        ledger=ledger.model,
        parent=general,
        slug=slug if existing is None else f"{slug}-{suffix}",
        code=code,
        name="اختلاف تراز واردات",
        account_class=general.account_class,
        normal_balance=general.normal_balance,
        sort_order=9999,
        is_active=True,
    )


def _live_import_journals(base_code, *, ledger):
    return (
        JournalEntry.objects.filter(ledger__code=ledger.id, document_code__startswith=base_code)
        .exclude(status_ref_id=JournalEntry.STATUS_VOID)
        .exclude(corrections__isnull=False)
        .exclude(corrects__isnull=False)
    )


def _claim_document_code(base_code, *, ledger, force, parsed):
    """کد سند جدید؛ None اگر همین دوره قبلاً وارد شده و جایگزینی خواسته نشده."""
    from logic.document_issuance import DocumentIssuanceService

    service = DocumentIssuanceService()
    for journal in _live_import_journals(base_code, ledger=ledger):
        if journal.status == JournalEntry.STATUS_POSTED:
            if not force:
                parsed.warnings.append(
                    f"سند {journal.document_code} قبلاً وارد و قطعی شده است؛ "
                    "برای جایگزینی گزینه «جایگزینی تراز قبلی» را فعال کنید."
                )
                return None
            service.issue_correction(journal, reason="جایگزینی با فایل اکسل جدید")
        else:
            service.retire_draft(journal, reason="جایگزینی با فایل اکسل جدید")
    used = JournalEntry.objects.filter(
        ledger__code=ledger.id, document_code__startswith=base_code,
    ).count()
    return base_code if used == 0 else f"{base_code}-R{used + 1}"


def _write_balance_journal(lines, *, code, entry_date, description, ledger, approve, accounts_by_path):
    from backend.models import JournalLine
    from logic.accounting import _allocate_document

    ledger_row, number, code = _allocate_document(
        ledger=ledger, entry_date=entry_date, document_code=code,
    )
    journal = JournalEntry.objects.create(
        ledger=ledger_row, document_number=number, document_code=code,
        entry_type="opening", entry_date=entry_date, description=description,
        status=JournalEntry.STATUS_DRAFT,
    )
    JournalLine.objects.bulk_create(
        [
            JournalLine(
                journal=journal,
                account=accounts_by_path[row.code] if row is not None else accounts_by_path[None],
                debit=debit, credit=credit, line_number=index,
                description=(row.name if row is not None else "اختلاف تراز فایل وارداتی")[:500],
            )
            for index, (row, debit, credit) in enumerate(lines, 1)
        ],
        batch_size=1000,
    )
    journal.refresh_totals(save=True)
    if approve:
        journal.post()
    return journal


def _import_trial_balances(parsed, *, ledger, approve, force):
    stats = {"opening_lines": 0, "turnover_lines": 0, "imbalance": 0, "journals_created": 0}
    date_from, date_to = parsed.metadata.get("date_from"), parsed.metadata.get("date_to")
    if not date_to:
        parsed.errors.append("تاریخ پایان دوره («تا تاریخ») در سربرگ تراز یافت نشد.")
        return stats

    leaves = _leaf_balance_rows(parsed)
    opening, turnover = _balance_lines(leaves)
    accounts_by_path = {
        account.path: account for account in ledger.accounts().filter(path__in={row.code for row in leaves})
    }
    parent_ids = set(ledger.accounts().exclude(parent__isnull=True).values_list("parent_id", flat=True))
    for row in leaves:
        account = accounts_by_path.get(row.code)
        if account is None:
            parsed.errors.append(f"حساب {row.code} در کدینگ یافت نشد.")
        elif account.id in parent_ids:
            parsed.errors.append(f"حساب {row.code} زیرحساب دارد و مانده مستقیم نمی‌پذیرد.")
    if parsed.errors:
        return stats

    opening_code, turnover_code = _balance_document_codes(parsed)
    jobs = (
        (opening, opening_code, parse_jalali_date(date_from or date_to), "سند افتتاحیه وارداتی"),
        (turnover, turnover_code, parse_jalali_date(date_to),
         f"گردش وارداتی {date_from or ''} تا {date_to}"),
    )
    for lines, code, entry_date, description in jobs:
        if not lines:
            continue
        base_code = code
        difference = _imbalance(lines)
        if difference:
            try:
                accounts_by_path[None] = _balancing_account(ledger)
            except ValueError as exc:
                parsed.errors.append(str(exc))
                return stats
            lines = lines + [(None, -difference, Decimal(0)) if difference < 0 else (None, Decimal(0), difference)]
            stats["imbalance"] += int(difference)
            parsed.warnings.append(
                f"{description}: فایل ناتراز است؛ اختلاف {int(abs(difference)):,} ریال "
                "به‌صورت خودکار در حساب فنی «اختلاف تراز واردات» ثبت شد."
            )
        code = _claim_document_code(base_code, ledger=ledger, force=force, parsed=parsed)
        if code is None:
            continue
        _write_balance_journal(
            lines, code=code, entry_date=entry_date, description=description,
            ledger=ledger, approve=approve, accounts_by_path=accounts_by_path,
        )
        stats["journals_created"] += 1
        stats["opening_lines" if base_code == opening_code else "turnover_lines"] = len(lines)
    return stats


def import_excel_file(file_obj, *, dry_run=False, approve=False, force=False, ledger=OFFICE_LEDGER):
    parsed = parse_excel_file(file_obj)
    parsed.recognized_accounts = _recognized_accounts(parsed, ledger=ledger)
    if parsed.errors:
        return _build_report(parsed, dry_run=dry_run, committed=False)

    if parsed.detail_ledger_rows:
        parsed.warnings.append(DETAIL_LEDGER_SAMPLE_NOTE)

    if dry_run:
        stats = _estimate_import_stats(parsed, ledger=ledger)
        balanced = not any(
            w for w in parsed.warnings
            if "گردش بدهکار" in w or "افتتاحیه بدهکار" in w
        )
        return _build_report(parsed, dry_run=True, committed=False, stats=stats, balanced=balanced)

    with transaction.atomic():
        with defer_account_closure_rebuild():
            stats = _import_chart_rows(parsed, ledger=ledger)
        stats.update(_import_trial_balances(parsed, ledger=ledger, approve=approve, force=force))
        if parsed.errors:
            transaction.set_rollback(True)

    if parsed.errors:
        return _build_report(parsed, dry_run=False, committed=False, stats=stats)

    balanced = not any(
        w for w in parsed.warnings
        if "گردش بدهکار" in w or "افتتاحیه بدهکار" in w
    )
    return _build_report(parsed, dry_run=False, committed=True, stats=stats, balanced=balanced)


def _build_report(parsed, *, dry_run, committed, stats=None, balanced=True):
    totals = parsed.metadata.get("general_totals")
    return {
        "dry_run": dry_run,
        "committed": committed,
        "balanced": balanced,
        "metadata": {
            "date_from": parsed.metadata.get("date_from"),
            "date_to": parsed.metadata.get("date_to"),
            "doc_from": parsed.metadata.get("doc_from"),
            "doc_to": parsed.metadata.get("doc_to"),
        },
        "counts": {
            "general_rows": len(parsed.general_rows),
            "subsidiary_rows": len(parsed.subsidiary_rows),
            "detailed_rows": len(parsed.detailed_rows),
            "detail_ledger_rows": len(parsed.detail_ledger_rows),
        },
        "stats": stats or {},
        "trial_totals": {
            "turnover_debit": int(totals.turnover_debit) if totals else 0,
            "turnover_credit": int(totals.turnover_credit) if totals else 0,
            "turnover_balanced": bool(totals and totals.turnover_debit == totals.turnover_credit),
        },
        "detail_ledger_account": parsed.detail_ledger_account_code,
        "recognized_accounts": parsed.recognized_accounts,
        "account_detection": {
            "total": len(parsed.recognized_accounts),
            "new": sum(1 for row in parsed.recognized_accounts if row["status"] == "new"),
            "existing": sum(1 for row in parsed.recognized_accounts if row["status"] == "existing"),
        },
        "import_mode": {
            "chart_from_trial_balance": True,
            "balances_from_trial_balance": True,
            "detail_sample_account": parsed.detail_ledger_account_code or "",
        },
        "warnings": parsed.warnings,
        "errors": parsed.errors,
    }
