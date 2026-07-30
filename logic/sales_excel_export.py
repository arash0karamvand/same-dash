"""خروجی اکسل فاکتور فروش — قالب سام اکسون (پر شدن خودکار از sale_to_dict)."""

from io import BytesIO
from pathlib import Path

from logic.jalali import date_to_jalali

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "sale_invoice_template.xlsx"
LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "company_logo.png"
LINE_ROW_START = 7
LINE_ROW_COUNT = 10


def _jalali_from_iso(iso_str):
    if not iso_str:
        return None
    from datetime import datetime

    from django.utils import timezone

    raw = str(iso_str).strip()
    if "T" in raw:
        dt = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        if timezone.is_aware(dt):
            dt = timezone.localtime(dt)
        d = dt.date()
    else:
        from datetime import date as date_cls

        d = date_cls.fromisoformat(raw[:10])
    jy, jm, jd = date_to_jalali(d)
    return f"{jy:04d}/{jm:02d}/{jd:02d}"


def _money_cell(value):
    if value is None:
        return None
    return int(value)


def _line_items(sale):
    items = list(sale.get("line_items") or [])
    if items:
        return items
    if sale.get("amounts_masked"):
        return []
    amount = sale.get("amount")
    if amount is None:
        amount = sale.get("final_amount")
    amount = int(amount or 0)
    desc = (sale.get("description") or "").strip() or "فروش"
    return [
        {
            "product_name": desc,
            "product_model": "",
            "fabric": "",
            "color_name": "",
            "quantity": 1,
            "unit_price": amount,
            "line_total": amount,
        }
    ]


def _write_cell(ws, row, col, value):
    """نوشتن در سلول؛ اگر داخل merge باشد، سلول بالا-چپ محدوده استفاده می‌شود."""
    from openpyxl.cell.cell import MergedCell

    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for merged in ws.merged_cells.ranges:
            if (
                merged.min_row <= row <= merged.max_row
                and merged.min_col <= col <= merged.max_col
            ):
                ws.cell(row=merged.min_row, column=merged.min_col, value=value)
                return
    cell.value = value


def _product_label(item):
    name = (item.get("product_name") or "").strip() or "—"
    specs = [item.get("fabric"), item.get("color_name"), item.get("product_model")]
    extra = " · ".join(x for x in specs if x)
    return f"{name} ({extra})" if extra else name


def _embed_logo(ws):
    if not LOGO_PATH.is_file():
        return
    try:
        from openpyxl.drawing.image import Image

        img = Image(str(LOGO_PATH))
        img.width = 220
        img.height = 70
        ws.add_image(img, "B1")
    except Exception:
        pass


def build_sale_excel_workbook(sale):
    from openpyxl import load_workbook

    if not TEMPLATE_PATH.is_file():
        raise FileNotFoundError(f"قالب فاکتور یافت نشد: {TEMPLATE_PATH}")

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    ws.sheet_view.rightToLeft = True
    _embed_logo(ws)

    customer = sale.get("customer_name") or ""
    phone = sale.get("customer_phone") or ""
    address = (sale.get("customer_address") or "").strip() or "—"
    sold = _jalali_from_iso(sale.get("sold_at"))

    ws["B4"] = f"نام مشتری : {customer}"
    ws["D4"] = f"شماره تماس: {phone}"
    _write_cell(ws, 4, 6, f" تاریخ فاکتور:{sold or ''}")
    ws["B5"] = f"آدرس: {address}"

    items = _line_items(sale)
    line_sum = 0
    for i in range(LINE_ROW_COUNT):
        row = LINE_ROW_START + i
        ws.cell(row=row, column=2, value=i + 1)
        if i < len(items):
            item = items[i]
            qty = int(item.get("quantity") or 0)
            unit = _money_cell(item.get("unit_price"))
            total = _money_cell(item.get("line_total"))
            if total is None and unit is not None:
                total = unit * qty
            _write_cell(ws, row, 3, _product_label(item))
            _write_cell(ws, row, 4, qty or None)
            _write_cell(ws, row, 5, unit)
            _write_cell(ws, row, 6, total if total is not None else 0)
            if total is not None:
                line_sum += total
        else:
            _write_cell(ws, row, 3, None)
            _write_cell(ws, row, 4, None)
            _write_cell(ws, row, 5, None)
            _write_cell(ws, row, 6, 0)

    if not sale.get("amounts_masked"):
        amount = _money_cell(sale.get("amount"))
        discount = _money_cell(sale.get("discount"))
        final = _money_cell(sale.get("final_amount"))
        paid = _money_cell(sale.get("paid_amount"))
        balance = _money_cell(sale.get("balance_due"))
        _write_cell(ws, 17, 6, amount if amount is not None else (line_sum or None))
        _write_cell(ws, 18, 6, discount if discount else None)
        _write_cell(ws, 19, 6, final)
        _write_cell(ws, 20, 6, paid if paid else None)
        _write_cell(ws, 21, 6, balance if balance is not None else None)
    else:
        for row_idx in (17, 18, 19, 20, 21):
            _write_cell(ws, row_idx, 6, None)

    delivery = _jalali_from_iso(sale.get("delivery_date"))
    _write_cell(ws, 22, 6, delivery if delivery and delivery != "—" else None)

    note_parts = []
    if sale.get("description"):
        note_parts.append(str(sale.get("description")).strip())
    if sale.get("accounting_mode_display"):
        note_parts.append(f"نوع ثبت حسابداری: {sale.get('accounting_mode_display')}")
    if sale.get("payment_method_display"):
        note_parts.append(f"روش پرداخت: {sale.get('payment_method_display')}")
    seller = sale.get("seller_name") or sale.get("recorded_by") or ""
    if seller:
        note_parts.append(f"کارشناس فروش: {seller}")
    _write_cell(ws, 19, 2, "\n".join(note_parts) if note_parts else None)

    return wb


def sale_excel_bytes(sale_dict):
    wb = build_sale_excel_workbook(sale_dict)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sale_excel_filename(sale_dict):
    """نام فایل — پیش‌فرض ASCII برای سازگاری با هدر HTTP."""
    sale_id = sale_dict.get("id") or "export"
    inv = (sale_dict.get("invoice_number") or "").strip()
    safe_inv = "".join(c if c.isalnum() and ord(c) < 128 else "_" for c in inv).strip("_")
    if safe_inv:
        return f"sale_{safe_inv}.xlsx"
    return f"sale_{sale_id}.xlsx"


def sale_excel_download_name(sale_dict):
    """نام پیشنهادی برای کاربر (ممکن است فارسی باشد)."""
    name = (sale_dict.get("customer_name") or "").strip()
    if name and name != "—":
        safe = "".join(c if c.isalnum() or c in " _-" else "_" for c in name)
        safe = safe.strip().replace(" ", "_")[:60]
        if safe.replace("_", ""):
            return f"{safe}.xlsx"
    return sale_excel_filename(sale_dict)


def content_disposition_attachment(filename, fallback_ascii=None):
    from urllib.parse import quote

    fallback = fallback_ascii or sale_excel_filename({"id": "export"})
    if filename and all(ord(c) < 128 for c in filename):
        return f'attachment; filename="{filename}"'
    quoted = quote(filename or fallback)
    return f'attachment; filename="{fallback}"; filename*=UTF-8\'\'{quoted}'
