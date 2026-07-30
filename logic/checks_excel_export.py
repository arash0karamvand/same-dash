"""خروجی اکسل فرم لیست چک‌های دریافتی."""

from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from logic.jalali import date_to_jalali

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "check_form_template.xlsx"
CHECK_FORM_MAX_ROWS = 16
DATA_ROW_START = 7
DATA_ROW_COUNT = CHECK_FORM_MAX_ROWS


def _jalali_date(value):
    if not value:
        return ""
    if hasattr(value, "isoformat"):
        d = value
    else:
        from datetime import date as date_cls

        d = date_cls.fromisoformat(str(value)[:10])
    jy, jm, jd = date_to_jalali(d)
    return f"{jy:04d}/{jm:02d}/{jd:02d}"


def _write_cell(ws, row, col, value):
    from openpyxl.cell.cell import MergedCell

    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for merged in ws.merged_cells.ranges:
            if (
                merged.min_row <= row <= merged.max_row
                and merged.min_col <= col <= merged.max_col
            ):
                ws.cell(row=merged.min_row, column=merged.min_col, value=value)
                return
    cell.value = value


def _row_from_installment(inst, idx):
    sale = inst.sale
    customer = sale.customer.full_name if sale and sale.customer_id else ""
    received = inst.received_at or getattr(inst, "created_at", None)
    if received and hasattr(received, "date"):
        received = received.date() if hasattr(received, "hour") else received
    seller = ""
    if sale and sale.seller_id:
        seller = sale.seller.full_name or ""
    return {
        "row": idx,
        "customer": customer,
        "received_at": _jalali_date(received),
        "bank_name": inst.bank_name or "",
        "due_date": _jalali_date(inst.due_date),
        "check_number": inst.check_number or "",
        "amount": int(inst.amount or 0),
        "receiver_name": inst.receiver_name or "",
        "seller_name": seller,
    }


def build_checks_excel_workbook(installments, *, seller_name="", manager_name=""):
    from openpyxl import load_workbook

    if not TEMPLATE_PATH.is_file():
        raise FileNotFoundError(f"قالب فرم چک یافت نشد: {TEMPLATE_PATH}")

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    ws.sheet_view.rightToLeft = True

    rows = [_row_from_installment(inst, i + 1) for i, inst in enumerate(installments[:DATA_ROW_COUNT])]
    for i in range(DATA_ROW_COUNT):
        row_num = DATA_ROW_START + i
        if i < len(rows):
            row = rows[i]
            _write_cell(ws, row_num, 2, row["row"])
            _write_cell(ws, row_num, 3, row["customer"])
            _write_cell(ws, row_num, 4, row["received_at"])
            _write_cell(ws, row_num, 5, row["bank_name"])
            _write_cell(ws, row_num, 6, row["due_date"])
            _write_cell(ws, row_num, 7, row["check_number"])
            _write_cell(ws, row_num, 8, row["amount"])
            _write_cell(ws, row_num, 9, row["receiver_name"])
        else:
            for col in range(2, 10):
                _write_cell(ws, row_num, col, None)

    expert = seller_name or (rows[0]["seller_name"] if rows else "")
    if expert:
        _write_cell(ws, 23, 2, f"کارشناس فروش : {expert}")
    if manager_name:
        _write_cell(ws, 23, 6, f"مسئول فروش : {manager_name}")

    return wb


def checks_excel_bytes(installments, **meta):
    wb = build_checks_excel_workbook(installments, **meta)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def checks_excel_filename(prefix="checks"):
    return f"{prefix}.xlsx"


def content_disposition_attachment(filename):
    fallback = checks_excel_filename()
    if filename and all(ord(c) < 128 for c in filename):
        return f'attachment; filename="{filename}"'
    quoted = quote(filename or fallback)
    return f'attachment; filename="{fallback}"; filename*=UTF-8\'\'{quoted}'
