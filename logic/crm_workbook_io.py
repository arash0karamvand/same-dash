"""خواندن و نوشتن فایل CRM (xlsb/xlsx) با همان تب‌ها و ستون‌ها."""

from __future__ import annotations

import io
import logging
import tempfile
from pathlib import Path
from typing import BinaryIO

import openpyxl
from pyxlsb import open_workbook

from logic.crm_workbook_schema import (
    CRM_COLUMNS,
    CRM_SHEETS,
    DEFAULT_WORKBOOK_PASSWORD,
    INVOICE_COLUMN,
    MULTI_ROW_SHEETS,
    SHEET_BY_EXCEL_NAME,
)

logger = logging.getLogger(__name__)


def _normalize_cell(value):
    if value is None:
        return None
    if isinstance(value, float) and value == int(value):
        return int(value)
    if isinstance(value, str):
        s = value.strip()
        return s if s else None
    return value


def _row_to_dict(sheet_key: str, values: list) -> dict:
    cols = CRM_COLUMNS[sheet_key]
    data = {}
    for i, col in enumerate(cols):
        if i < len(values):
            data[col] = _normalize_cell(values[i])
        else:
            data[col] = None
    return data


def _is_empty_row(data: dict) -> bool:
    return not any(v not in (None, "", "-", 0, 0.0) for v in data.values())


def _invoice_from_row(sheet_key: str, data: dict) -> str:
    col = INVOICE_COLUMN.get(sheet_key, "")
    raw = data.get(col)
    if raw is None:
        return ""
    return str(raw).strip()


def _decrypt_to_buffer(file_obj: BinaryIO, password: str | None) -> io.BytesIO:
    import msoffcrypto

    buf = io.BytesIO()
    if hasattr(file_obj, "seek"):
        file_obj.seek(0)
    raw = file_obj.read()
    src = io.BytesIO(raw)
    office = msoffcrypto.OfficeFile(src)
    if office.is_encrypted():
        pwd = password or DEFAULT_WORKBOOK_PASSWORD
        office.load_key(password=pwd)
        office.decrypt(buf)
    else:
        buf.write(raw)
    buf.seek(0)
    return buf


def read_xlsb_workbook(file_obj: BinaryIO, *, password: str | None = None) -> dict[str, list[dict]]:
    """برگرداندن {sheet_key: [row_data, ...]}."""
    buf = _decrypt_to_buffer(file_obj, password)
    parsed: dict[str, list[dict]] = {k: [] for k in CRM_SHEETS}

    with open_workbook(buf) as wb:
        for excel_name in wb.sheets:
            sheet_key = SHEET_BY_EXCEL_NAME.get(excel_name)
            if not sheet_key:
                continue
            rows_out: list[dict] = []
            with wb.get_sheet(excel_name) as sh:
                for i, row in enumerate(sh.rows()):
                    vals = [c.v for c in row]
                    if i == 0:
                        continue  # header
                    if sheet_key == "financial" and i == 1:
                        continue  # sub-header واریزی‌ها
                    data = _row_to_dict(sheet_key, vals)
                    if _is_empty_row(data):
                        continue
                    rows_out.append(data)
            parsed[sheet_key] = rows_out
    return parsed


def read_xlsx_workbook(path: str | Path) -> dict[str, list[dict]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parsed: dict[str, list[dict]] = {k: [] for k in CRM_SHEETS}
    try:
        for sheet_key, excel_name in CRM_SHEETS.items():
            if excel_name not in wb.sheetnames:
                continue
            ws = wb[excel_name]
            rows_out: list[dict] = []
            for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                if i == 1:
                    continue
                if sheet_key == "financial" and i == 2:
                    continue
                data = _row_to_dict(sheet_key, list(row))
                if _is_empty_row(data):
                    continue
                rows_out.append(data)
            parsed[sheet_key] = rows_out
    finally:
        wb.close()
    return parsed


def read_uploaded_workbook(upload, *, password: str | None = None) -> dict[str, list[dict]]:
    name = (upload.name or "").lower()
    if name.endswith(".xlsb"):
        return read_xlsb_workbook(upload, password=password)
    if name.endswith(".xlsx"):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            for chunk in upload.chunks():
                tmp.write(chunk)
            tmp_path = tmp.name
        try:
            return read_xlsx_workbook(tmp_path)
        finally:
            Path(tmp_path).unlink(missing_ok=True)
    raise ValueError("فقط فایل .xlsb یا .xlsx پذیرفته می‌شود")


def export_workbook_xlsx(rows_by_sheet: dict[str, list[dict]]) -> bytes:
    """خروجی xlsx با همان نام برگه‌ها و ستون‌ها."""
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    for sheet_key, excel_name in CRM_SHEETS.items():
        ws = wb.create_sheet(title=excel_name[:31])
        cols = CRM_COLUMNS[sheet_key]
        ws.append(cols)
        if sheet_key == "financial":
            sub = [None, None, "تاریخ واریز", "مبلغ واریز شده", "شماره حساب", "سایر"]
            ws.append(sub)
        for row in rows_by_sheet.get(sheet_key) or []:
            ws.append([row.get(c) for c in cols])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def assign_sort_orders(sheet_key: str, rows: list[dict]) -> list[tuple[str, int, dict]]:
    """(invoice_ref, sort_order, data) برای ذخیره."""
    invoice_counts: dict[str, int] = {}
    result = []
    for idx, data in enumerate(rows):
        inv = _invoice_from_row(sheet_key, data)
        if sheet_key in MULTI_ROW_SHEETS:
            key = inv or f"__row_{idx}"
            n = invoice_counts.get(key, 0)
            invoice_counts[key] = n + 1
            sort_order = n
        elif not inv:
            sort_order = idx
        else:
            sort_order = 0
        result.append((inv, sort_order, data))
    return result
