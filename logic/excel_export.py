"""سیستم خروجی اکسل فرمت‌دار برای فرم‌های اداری، فاکتورها و گزارشات"""

import tempfile
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from django.conf import settings
from django.http import HttpResponse


class FormattedExcelExporter:
    """کلاس پایه برای خروجی اکسل با فرمت و طراحی حرفه‌ای"""

    def __init__(self, title):
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = title[:31]  # Excel sheet title limit
        self.current_row = 1

    def add_logo(self):
        """اضافه کردن لوگوی شرکت در گوشه بالا"""
        logo_path = Path(settings.BASE_DIR) / "ui" / "dist" / "company_logo.png"
        if logo_path.exists():
            try:
                from openpyxl.drawing.image import Image

                img = Image(str(logo_path))
                img.height = 60
                img.width = 120
                self.ws.add_image(img, "A1")
            except Exception:
                pass  # If logo fails, continue without it
        self.current_row = 4

    def add_header(self, text, code=""):
        """هدر فرم با کد مدرک"""
        cell = self.ws[f"A{self.current_row}"]
        cell.value = text
        cell.font = Font(name="B Nazanin", size=16, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

        if code:
            code_cell = self.ws[f"G{self.current_row}"]
            code_cell.value = f"کد مدرک: {code}"
            code_cell.font = Font(name="B Nazanin", size=10, bold=True)
            code_cell.alignment = Alignment(horizontal="left")

        self.current_row += 2

    def add_field(self, label, value, row=None, col_label="A", col_value="C"):
        """فیلد label: value با فرمت"""
        r = row or self.current_row

        label_cell = self.ws[f"{col_label}{r}"]
        label_cell.value = label
        label_cell.font = Font(name="B Nazanin", size=11, bold=True)
        label_cell.alignment = Alignment(horizontal="right")

        value_cell = self.ws[f"{col_value}{r}"]
        value_cell.value = value
        value_cell.font = Font(name="B Nazanin", size=11)
        value_cell.alignment = Alignment(horizontal="right")

        if not row:
            self.current_row += 1

    def add_table_header(self, headers, start_row=None, start_col=1):
        """هدر جدول با رنگ و فرمت"""
        r = start_row or self.current_row

        for i, header in enumerate(headers):
            col = start_col + i
            cell = self.ws.cell(r, col)
            cell.value = header
            cell.font = Font(name="B Nazanin", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="4472C4", end_color="4472C4", fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

        if not start_row:
            self.current_row += 1

    def add_table_row(self, values, row=None, start_col=1, is_total=False):
        """ردیف جدول"""
        r = row or self.current_row

        for i, value in enumerate(values):
            col = start_col + i
            cell = self.ws.cell(r, col)
            cell.value = value
            cell.font = Font(
                name="B Nazanin", size=10, bold=is_total
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            if is_total:
                cell.fill = PatternFill(
                    start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
                )

        if not row:
            self.current_row += 1

    def add_signature_row(self, labels):
        """ردیف امضاها"""
        self.current_row += 2

        for i, label in enumerate(labels):
            col_offset = i * 3
            col = chr(65 + col_offset)  # A, D, G, ...

            label_cell = self.ws[f"{col}{self.current_row}"]
            label_cell.value = label
            label_cell.font = Font(name="B Nazanin", size=11, bold=True)
            label_cell.alignment = Alignment(horizontal="center")

            sign_cell = self.ws[f"{col}{self.current_row + 1}"]
            sign_cell.value = "تاریخ و امضا:"
            sign_cell.font = Font(name="B Nazanin", size=9)
            sign_cell.alignment = Alignment(horizontal="center")

        self.current_row += 3

    def merge_cells(self, start, end):
        """ادغام سلول‌ها"""
        self.ws.merge_cells(f"{start}:{end}")

    def set_column_width(self, column, width):
        """تنظیم عرض ستون"""
        self.ws.column_dimensions[column].width = width

    def save_response(self, filename):
        """ذخیره و برگرداندن HttpResponse"""
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        self.wb.save(response)
        return response

    def save_to_temp(self):
        """ذخیره در فایل موقت"""
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        self.wb.save(tmp.name)
        return tmp.name


# ========== توابع خروجی اکسل برای فرم‌های مختلف ==========


def export_assistance_request_excel(obj):
    """خروجی اکسل درخواست مساعده"""
    exp = FormattedExcelExporter("درخواست مساعده")
    exp.add_logo()
    exp.add_header("درخواست مساعده", code="210/FR/09/00")

    exp.add_field("نام و نام خانوادگی:", obj.employee.get_full_name())
    exp.add_field("کد ملی:", obj.national_code or "—")
    exp.add_field("واحد:", obj.department or "—")
    exp.add_field("مبلغ درخواستی:", f"{obj.amount:,} ریال")
    exp.add_field("علت درخواست:", obj.reason or "—")
    exp.add_field("تاریخ:", obj.request_date.strftime("%Y/%m/%d"))

    exp.current_row += 1
    if obj.notes:
        exp.add_field("توضیحات:", obj.notes)

    exp.add_signature_row(["درخواست کننده", "تایید مدیر واحد", "مدیر عامل"])

    return exp.save_response(f"assistance_{obj.reference_code}.xlsx")


def export_petty_cash_excel(obj):
    """خروجی اکسل تنخواه"""
    exp = FormattedExcelExporter("تنخواه")
    exp.add_logo()
    exp.add_header("فرم درخواست تنخواه", code="220/FR/09/00")

    exp.add_field("درخواست کننده:", obj.requester.get_full_name())
    exp.add_field("مبلغ:", f"{obj.amount:,} ریال")
    exp.add_field("منظور:", obj.purpose)
    exp.add_field("تاریخ درخواست:", obj.request_date.strftime("%Y/%m/%d"))

    if obj.status in (obj.STATUS_PAID, obj.STATUS_SETTLED):
        exp.current_row += 1
        exp.add_field("تاریخ پرداخت:", obj.payment_date.strftime("%Y/%m/%d") if obj.payment_date else "—")

    if obj.status == obj.STATUS_SETTLED:
        exp.add_field("مبلغ تسویه:", f"{obj.settlement_amount:,} ریال" if obj.settlement_amount else "—")
        exp.add_field("تاریخ تسویه:", obj.settlement_date.strftime("%Y/%m/%d") if obj.settlement_date else "—")

    exp.add_signature_row(["درخواست کننده", "تایید مدیر", "امور مالی"])

    return exp.save_response(f"petty_cash_{obj.reference_code}.xlsx")


def export_warehouse_transfer_excel(obj):
    """خروجی اکسل جابجایی انبار"""
    exp = FormattedExcelExporter("جابجایی انبار")
    exp.add_logo()
    exp.add_header("فرم جابجایی بین انبار", code="300/FR/04/00")

    exp.add_field("از انبار:", obj.from_warehouse.name)
    exp.add_field("به انبار:", obj.to_warehouse.name)
    exp.add_field("تاریخ جابجایی:", obj.transfer_date.strftime("%Y/%m/%d"))
    exp.add_field("تعداد نفرات مورد نیاز:", str(obj.required_personnel))

    # جدول آیتم‌ها
    exp.current_row += 2
    exp.add_table_header(["ردیف", "کد کالا", "شرح", "تعداد/مقدار", "واحد", "توضیحات"])

    for line in obj.lines.all():
        exp.add_table_row([
            line.line_number,
            line.product_code,
            line.product_description,
            line.quantity,
            line.unit,
            line.notes or "—",
        ])

    # کنترل کیفیت
    exp.current_row += 2
    exp.add_field("آیا محصول بدون عیب تحویل داده شده است؟", "بله" if not obj.has_defects else "خیر")
    if obj.has_defects:
        exp.add_field("شرح عیوب:", obj.defect_description)

    exp.add_signature_row(["مسئول انبار", "کارشناس فروش", "کنترل کیفیت"])

    return exp.save_response(f"transfer_{obj.reference_code}.xlsx")


def export_production_order_excel(obj):
    """خروجی اکسل سفارش تولید"""
    exp = FormattedExcelExporter("سفارش تولید")
    exp.add_logo()
    exp.add_header("فرم سفارش تولید / خدمات", code="500/FR/06/00")

    exp.add_field("نام نمایندگی / مشتری:", obj.customer_name or "—")
    exp.add_field("شماره فاکتور:", obj.invoice_number or "—")
    exp.add_field("شماره سفارش تولید:", obj.reference_code)
    exp.add_field("مدل / نام کالا:", obj.product_model or "—")
    exp.add_field("تاریخ ثبت سفارش:", obj.order_date.strftime("%Y/%m/%d"))
    exp.add_field("تاریخ تحویل:", obj.delivery_date.strftime("%Y/%m/%d"))
    
    priority_label = dict(obj.PRIORITY_CHOICES).get(obj.priority, "—")
    exp.add_field("اولویت تولید:", priority_label)

    # جدول آیتم‌ها
    exp.current_row += 2
    exp.add_table_header([
        "ردیف",
        "تعداد",
        "رنگ چوب",
        "رنگ صفحات",
        "رنگ پارچه",
        "کالیته",
        "کد پارچه",
    ])

    for line in obj.lines.all():
        exp.add_table_row([
            line.line_number,
            line.quantity,
            line.wood_color or "—",
            line.panel_color or "—",
            line.fabric_color or "—",
            line.fabric_quality or "—",
            line.fabric_code or "—",
        ])

    if obj.cushion_notes:
        exp.current_row += 1
        exp.add_field("کوسن‌ها:", obj.cushion_notes)

    if obj.general_notes:
        exp.add_field("توضیحات:", obj.general_notes)

    exp.add_signature_row(["تنظیم کننده", "کارشناس فروش", "مدیریت", "برنامه‌ریزی تولید"])

    return exp.save_response(f"production_{obj.reference_code}.xlsx")


def export_attendance_confirmation_excel(obj):
    """خروجی اکسل تایید کارکرد"""
    exp = FormattedExcelExporter("تایید کارکرد")
    exp.add_logo()
    exp.add_header("تایید کارکرد ماهانه", code="410/FR/HR/00")

    exp.add_field("نام و نام خانوادگی:", obj.employee.get_full_name())
    exp.add_field("ماه:", obj.month)
    exp.add_field("تعداد روزهای کاری:", str(obj.days_worked))
    exp.add_field("ساعات اضافه‌کاری:", f"{obj.overtime_hours} ساعت")
    exp.add_field("روزهای غیبت:", str(obj.absence_days))
    exp.add_field("روزهای مرخصی:", str(obj.leave_days))

    if obj.notes:
        exp.current_row += 1
        exp.add_field("توضیحات:", obj.notes)

    exp.add_signature_row(["کارمند", "سرپرست", "امور اداری"])

    return exp.save_response(f"attendance_{obj.reference_code}.xlsx")


def export_sale_invoice_excel(sale):
    """فاکتور فروش با فرمت کامل و حرفه‌ای"""
    exp = FormattedExcelExporter("فاکتور فروش")
    exp.add_logo()

    # Header با رنگ
    exp.ws["A3"] = "فاکتور فروش"
    exp.ws["A3"].font = Font(name="B Nazanin", size=18, bold=True, color="FFFFFF")
    exp.ws["A3"].fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    exp.ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    exp.merge_cells("A3", "G3")

    exp.ws["H3"] = f"شماره: {sale.reference_code}"
    exp.ws["H3"].font = Font(name="B Nazanin", size=12, bold=True)
    exp.ws["H3"].alignment = Alignment(horizontal="left")

    # اطلاعات مشتری
    exp.current_row = 5
    exp.add_field("نام مشتری:", sale.customer.full_name if sale.customer else "نقدی")
    if sale.customer:
        exp.add_field("تلفن:", sale.customer.phone)
        if sale.customer.address:
            exp.add_field("آدرس:", sale.customer.address)
    exp.add_field("تاریخ:", sale.sale_date.strftime("%Y/%m/%d"))
    exp.add_field("فروشنده:", sale.seller.user.get_full_name() if sale.seller else "—")

    # جدول آیتم‌ها
    exp.current_row += 1
    headers = ["ردیف", "شرح کالا", "تعداد", "قیمت واحد", "جمع"]
    exp.add_table_header(headers)

    line_items = sale.line_items.all() if hasattr(sale, 'line_items') else []
    for idx, item in enumerate(line_items, 1):
        exp.add_table_row([
            idx,
            item.product_name,
            item.quantity,
            f"{item.unit_price:,}",
            f"{item.total_price:,}",
        ])

    # مبالغ نهایی
    exp.current_row += 1
    exp.add_field("جمع:", f"{sale.subtotal:,} ریال", col_label="D", col_value="E")
    if sale.discount_amount and sale.discount_amount > 0:
        exp.add_field("تخفیف:", f"{sale.discount_amount:,} ریال", col_label="D", col_value="E")
    exp.add_field("مبلغ نهایی:", f"{sale.final_amount:,} ریال", col_label="D", col_value="E")
    exp.add_field("مبلغ پرداخت شده:", f"{sale.paid_amount:,} ریال", col_label="D", col_value="E")
    
    remaining = sale.final_amount - sale.paid_amount
    if remaining > 0:
        exp.add_field("باقیمانده:", f"{remaining:,} ریال", col_label="D", col_value="E")

    # امضا
    exp.add_signature_row(["فروشنده", "مشتری", "مدیر"])

    # تنظیم عرض ستون‌ها
    exp.set_column_width("A", 8)
    exp.set_column_width("B", 30)
    exp.set_column_width("C", 10)
    exp.set_column_width("D", 15)
    exp.set_column_width("E", 15)

    return exp.save_response(f"invoice_{sale.reference_code}.xlsx")


def export_check_info_excel(installment):
    """خروجی اکسل اطلاعات چک (بهبود یافته)"""
    exp = FormattedExcelExporter("اطلاعات چک")
    exp.add_logo()
    exp.add_header("فرم اطلاعات چک", code="320/FR/ACC/00")

    exp.add_field("شماره چک:", installment.check_number or "—")
    exp.add_field("مبلغ:", f"{installment.amount:,} ریال")
    exp.add_field("تاریخ سررسید:", installment.due_date.strftime("%Y/%m/%d") if installment.due_date else "—")
    
    if hasattr(installment, 'sale') and installment.sale:
        exp.add_field("فاکتور مربوطه:", installment.sale.reference_code)
        if installment.sale.customer:
            exp.add_field("نام مشتری:", installment.sale.customer.full_name)

    exp.add_field("بانک:", installment.bank_name or "—")
    exp.add_field("شعبه:", installment.branch_name or "—")
    exp.add_field("صاحب حساب:", installment.account_holder or "—")

    if installment.notes:
        exp.current_row += 1
        exp.add_field("توضیحات:", installment.notes)

    exp.add_signature_row(["تهیه کننده", "مدیر مالی", "تایید"])

    return exp.save_response(f"check_{installment.id}.xlsx")


def export_payment_order_excel(payment_data):
    """خروجی اکسل دستور پرداخت"""
    exp = FormattedExcelExporter("دستور پرداخت")
    exp.add_logo()
    exp.add_header("دستور پرداخت", code="330/FR/PAY/00")

    exp.add_field("گیرنده وجه:", payment_data.get("payee", "—"))
    exp.add_field("مبلغ:", f"{payment_data.get('amount', 0):,} ریال")
    exp.add_field("بابت:", payment_data.get("description", "—"))
    exp.add_field("تاریخ:", payment_data.get("date", "—"))

    payment_method = payment_data.get("method", "cash")
    method_label = {"cash": "نقدی", "check": "چک", "transfer": "انتقال بانکی"}.get(
        payment_method, "—"
    )
    exp.add_field("نحوه پرداخت:", method_label)

    if payment_method == "check":
        exp.add_field("شماره چک:", payment_data.get("check_number", "—"))
        exp.add_field("بانک:", payment_data.get("bank", "—"))

    exp.add_signature_row(["درخواست کننده", "تایید مالی", "پرداخت"])

    filename = payment_data.get("filename", "payment_order.xlsx")
    return exp.save_response(filename)
